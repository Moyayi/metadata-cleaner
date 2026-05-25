#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MetaCleaner v1.1 — Limpiador de Metadatos (seleccion multiple)
"""

# ── Importaciones estándar ──────────────────────────────────────────────
import os
import json
import threading
from pathlib import Path

# ── Interfaz gráfica ────────────────────────────────────────────────────
import customtkinter as ctk
from tkinter import filedialog

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD as _TkDnD
    _DND_AVAILABLE = True
except ImportError:
    _DND_AVAILABLE = False

# ── Procesamiento de documentos ─────────────────────────────────────────
from docx import Document
import openpyxl
from PIL import Image
from PIL.ExifTags import TAGS, GPSTAGS

try:
    from pypdf import PdfReader, PdfWriter
    _PDF_LIB = "pypdf"
except ImportError:
    try:
        import PyPDF2
        _PDF_LIB = "pypdf2"
    except ImportError:
        _PDF_LIB = None


# ══════════════════════════════════════════════════════════════════════════
#  PALETA DE COLORES
# ══════════════════════════════════════════════════════════════════════════
C = {
    "bg_main":    "#0D0D0D",
    "bg_panel":   "#161616",
    "bg_card":    "#1F1F1F",
    "bg_row_alt": "#181818",
    "bg_sel":     "#1A1735",
    "border":     "#2A2A2A",
    "accent":     "#7C6AF7",
    "accent_dim": "#5F50D4",
    "success":    "#22C55E",
    "success_dim":"#15803D",
    "error":      "#EF4444",
    "warning":    "#F59E0B",
    "txt_main":   "#EFEFEF",
    "txt_dim":    "#6B7280",
    "txt_accent": "#A5A0FF",
    "table_head": "#1A1735",
}


# ══════════════════════════════════════════════════════════════════════════
#  BACKEND — Extracción
# ══════════════════════════════════════════════════════════════════════════

class MetadataExtractor:
    FORMATOS = {".json", ".docx", ".xlsx", ".pdf", ".png", ".jpg", ".jpeg"}

    @staticmethod
    def extraer(ruta: str) -> dict:
        ext = Path(ruta).suffix.lower()
        mapa = {
            ".docx": MetadataExtractor._docx,
            ".xlsx": MetadataExtractor._xlsx,
            ".pdf":  MetadataExtractor._pdf,
            ".png":  MetadataExtractor._imagen,
            ".jpg":  MetadataExtractor._imagen,
            ".jpeg": MetadataExtractor._imagen,
            ".json": MetadataExtractor._json,
        }
        if ext not in mapa:
            raise ValueError(f"Formato '{ext}' no soportado.")
        return mapa[ext](ruta)

    @staticmethod
    def _docx(ruta: str) -> dict:
        doc = Document(ruta)
        p = doc.core_properties
        return {
            "Título":             p.title or "",
            "Autor":              p.author or "",
            "Último editor":      p.last_modified_by or "",
            "Fecha creación":     str(p.created) if p.created else "",
            "Fecha modificación": str(p.modified) if p.modified else "",
            "Asunto":             p.subject or "",
            "Palabras clave":     p.keywords or "",
            "Descripción":        p.description or "",
            "Categoría":          p.category or "",
            "Revisión":           str(p.revision) if p.revision else "",
        }

    @staticmethod
    def _xlsx(ruta: str) -> dict:
        wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
        p = wb.properties
        return {
            "Título":             p.title or "",
            "Creador":            p.creator or "",
            "Último editor":      p.lastModifiedBy or "",
            "Fecha creación":     str(p.created) if p.created else "",
            "Fecha modificación": str(p.modified) if p.modified else "",
            "Asunto":             p.subject or "",
            "Descripción":        p.description or "",
            "Palabras clave":     p.keywords or "",
            "Categoría":          p.category or "",
            "Empresa":            p.company or "",
        }

    @staticmethod
    def _pdf(ruta: str) -> dict:
        if _PDF_LIB == "pypdf":
            reader = PdfReader(ruta)
            meta = reader.metadata or {}
            return {
                "Título":             _safe(meta.get("/Title")),
                "Autor":              _safe(meta.get("/Author")),
                "Creador":            _safe(meta.get("/Creator")),
                "Productor":          _safe(meta.get("/Producer")),
                "Asunto":             _safe(meta.get("/Subject")),
                "Palabras clave":     _safe(meta.get("/Keywords")),
                "Fecha creación":     _safe(meta.get("/CreationDate")),
                "Fecha modificación": _safe(meta.get("/ModDate")),
                "Número de páginas":  str(len(reader.pages)),
            }
        elif _PDF_LIB == "pypdf2":
            with open(ruta, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                meta = reader.metadata or {}
                return {
                    "Título":            _safe(meta.get("/Title")),
                    "Autor":             _safe(meta.get("/Author")),
                    "Creador":           _safe(meta.get("/Creator")),
                    "Productor":         _safe(meta.get("/Producer")),
                    "Asunto":            _safe(meta.get("/Subject")),
                    "Fecha creación":    _safe(meta.get("/CreationDate")),
                    "Número de páginas": str(len(reader.pages)),
                }
        else:
            raise RuntimeError("Instala: pip install pypdf")

    @staticmethod
    def _imagen(ruta: str) -> dict:
        img = Image.open(ruta)
        meta = {
            "Formato":       img.format or Path(ruta).suffix.upper().lstrip("."),
            "Modo de color": img.mode,
            "Dimensiones":   f"{img.width} x {img.height} px",
        }
        exif_raw = None
        if hasattr(img, "_getexif"):
            try:
                exif_raw = img._getexif()
            except Exception:
                pass
        if exif_raw:
            for tag_id, valor in exif_raw.items():
                nombre_tag = TAGS.get(tag_id, str(tag_id))
                if nombre_tag == "GPSInfo" and isinstance(valor, dict):
                    for gps_id, gps_val in valor.items():
                        meta[f"GPS / {GPSTAGS.get(gps_id, str(gps_id))}"] = _normalizar_exif(gps_val)
                else:
                    meta[nombre_tag] = _normalizar_exif(valor)
        return meta

    @staticmethod
    def _json(ruta: str) -> dict:
        tamano = os.path.getsize(ruta)
        with open(ruta, "r", encoding="utf-8", errors="replace") as f:
            datos = json.load(f)
        meta = {"Tamaño": f"{tamano:,} bytes", "Codificación": "UTF-8"}
        if isinstance(datos, dict):
            meta["Tipo raíz"]          = "Objeto {}"
            meta["Número de claves"]   = str(len(datos))
            meta["Claves principales"] = ", ".join(str(k) for k in list(datos.keys())[:8])
        elif isinstance(datos, list):
            meta["Tipo raíz"]           = "Array []"
            meta["Número de elementos"] = str(len(datos))
        return meta


# ══════════════════════════════════════════════════════════════════════════
#  BACKEND — Limpieza
# ══════════════════════════════════════════════════════════════════════════

class MetadataCleaner:
    @staticmethod
    def limpiar(ruta_entrada: str, ruta_salida: str) -> None:
        ext = Path(ruta_entrada).suffix.lower()
        mapa = {
            ".docx": MetadataCleaner._docx,
            ".xlsx": MetadataCleaner._xlsx,
            ".pdf":  MetadataCleaner._pdf,
            ".png":  MetadataCleaner._imagen,
            ".jpg":  MetadataCleaner._imagen,
            ".jpeg": MetadataCleaner._imagen,
            ".json": MetadataCleaner._json,
        }
        if ext not in mapa:
            raise ValueError(f"Formato '{ext}' no soportado.")
        mapa[ext](ruta_entrada, ruta_salida)

    @staticmethod
    def _docx(entrada: str, salida: str) -> None:
        doc = Document(entrada)
        p = doc.core_properties
        p.author = p.title = p.subject = p.keywords = ""
        p.description = p.last_modified_by = p.category = p.content_status = ""
        p.revision = 1
        doc.save(salida)

    @staticmethod
    def _xlsx(entrada: str, salida: str) -> None:
        wb = openpyxl.load_workbook(entrada)
        p = wb.properties
        p.creator = p.title = p.subject = p.description = ""
        p.keywords = p.lastModifiedBy = p.category = p.company = ""
        p.created = p.modified = None
        wb.save(salida)

    @staticmethod
    def _pdf(entrada: str, salida: str) -> None:
        vaciar = {"/Title": "", "/Author": "", "/Subject": "",
                  "/Creator": "", "/Producer": "", "/Keywords": ""}
        if _PDF_LIB == "pypdf":
            reader = PdfReader(entrada)
            writer = PdfWriter()
            for page in reader.pages:
                writer.add_page(page)
            writer.add_metadata(vaciar)
            with open(salida, "wb") as f:
                writer.write(f)
        elif _PDF_LIB == "pypdf2":
            with open(entrada, "rb") as fi:
                reader = PyPDF2.PdfReader(fi)
                writer = PyPDF2.PdfWriter()
                for page in reader.pages:
                    writer.add_page(page)
                writer.add_metadata(vaciar)
                with open(salida, "wb") as fo:
                    writer.write(fo)
        else:
            raise RuntimeError("No hay librería PDF disponible.")

    @staticmethod
    def _imagen(entrada: str, salida: str) -> None:
        with Image.open(entrada) as img:
            pixeles = img.tobytes()
            limpia = Image.frombytes(img.mode, img.size, pixeles)
            ext = Path(salida).suffix.lower()
            if ext in (".jpg", ".jpeg"):
                if limpia.mode not in ("RGB", "L"):
                    limpia = limpia.convert("RGB")
                limpia.save(salida, format="JPEG", quality=95, subsampling=0)
            else:
                limpia.save(salida, format="PNG", optimize=True)

    @staticmethod
    def _json(entrada: str, salida: str) -> None:
        with open(entrada, "r", encoding="utf-8", errors="replace") as f:
            datos = json.load(f)
        with open(salida, "w", encoding="utf-8") as f:
            json.dump(datos, f, ensure_ascii=False, indent=2)


# ══════════════════════════════════════════════════════════════════════════
#  HELPERS
# ══════════════════════════════════════════════════════════════════════════

def _safe(v) -> str:
    return "" if v is None else str(v)[:300]

def _normalizar_exif(v) -> str:
    if isinstance(v, bytes):
        return v.decode("utf-8", errors="replace").rstrip("\x00")
    return str(v)[:200]

def _fmt_tam(n: int) -> str:
    if n < 1024:       return f"{n} B"
    if n < 1024**2:    return f"{n/1024:.1f} KB"
    return             f"{n/1024**2:.1f} MB"

def _parsear_rutas_dnd(data: str) -> list:
    """Parsea el string de drop que puede contener una o varias rutas."""
    rutas, data = [], data.strip()
    while data:
        if data.startswith("{"):
            end = data.find("}")
            if end != -1:
                rutas.append(data[1:end])
                data = data[end + 1:].strip()
            else:
                rutas.append(data[1:])
                break
        else:
            partes = data.split(" ", 1)
            rutas.append(partes[0])
            data = partes[1].strip() if len(partes) > 1 else ""
    return [r for r in rutas if r]


# ══════════════════════════════════════════════════════════════════════════
#  UI — Tabla de metadatos
# ══════════════════════════════════════════════════════════════════════════

class TablaMetadatos(ctk.CTkScrollableFrame):
    def __init__(self, padre, **kwargs):
        super().__init__(padre, **kwargs)
        self._placeholder()

    def _placeholder(self):
        ctk.CTkLabel(
            self, text="Selecciona un archivo de la lista\npara ver sus metadatos",
            font=ctk.CTkFont(size=14), text_color=C["txt_dim"], justify="center",
        ).pack(expand=True, pady=50)

    def actualizar(self, meta: dict):
        for w in self.winfo_children():
            w.destroy()
        if not meta:
            ctk.CTkLabel(self, text="Sin metadatos detectados.",
                         text_color=C["txt_dim"]).pack(pady=20)
            return
        cab = ctk.CTkFrame(self, fg_color=C["table_head"], corner_radius=8)
        cab.pack(fill="x", pady=(0, 4))
        ctk.CTkLabel(cab, text="Campo", font=ctk.CTkFont(size=12, weight="bold"),
                     text_color=C["txt_accent"], width=170, anchor="w"
                     ).pack(side="left", padx=12, pady=8)
        ctk.CTkLabel(cab, text="Valor", font=ctk.CTkFont(size=12, weight="bold"),
                     text_color=C["txt_accent"], anchor="w"
                     ).pack(side="left", padx=8, pady=8)
        for i, (k, v) in enumerate(meta.items()):
            bg = C["bg_card"] if i % 2 == 0 else C["bg_row_alt"]
            f = ctk.CTkFrame(self, fg_color=bg, corner_radius=6)
            f.pack(fill="x", pady=1)
            ctk.CTkLabel(f, text=str(k), font=ctk.CTkFont(size=12),
                         text_color=C["txt_dim"], width=170, anchor="w"
                         ).pack(side="left", padx=12, pady=6)
            disp  = str(v) if v else "—"
            color = C["txt_main"] if v else C["txt_dim"]
            ctk.CTkLabel(f, text=disp, font=ctk.CTkFont(size=12),
                         text_color=color, anchor="w", wraplength=360
                         ).pack(side="left", padx=8, pady=6, fill="x", expand=True)

    def limpiar(self):
        for w in self.winfo_children():
            w.destroy()
        self._placeholder()


# ══════════════════════════════════════════════════════════════════════════
#  UI — Fila de archivo en la lista
# ══════════════════════════════════════════════════════════════════════════

# (icono, color, texto de estado)
_INFO_ESTADO = {
    "extrayendo": ("...", C["warning"],  "Leyendo metadatos"),
    "listo":      ("●",   C["accent"],   "Listo para limpiar"),
    "limpiando":  ("...", C["warning"],  "Limpiando"),
    "limpio":     ("✓",   C["success"],  "Limpio"),
    "error":      ("✗",   C["error"],    "Error"),
}

class FilaArchivo(ctk.CTkFrame):
    def __init__(self, padre, item: dict, cb_seleccionar, cb_eliminar, **kwargs):
        super().__init__(padre, **kwargs)
        self._item = item
        self._cb_sel = cb_seleccionar
        self._cb_del = cb_eliminar
        self._construir()

    def _construir(self):
        self.configure(fg_color=C["bg_card"], corner_radius=8)

        # Icono de estado
        ico, col, _ = _INFO_ESTADO.get(self._item["estado"], ("?", C["txt_dim"], ""))
        self._lbl_ico = ctk.CTkLabel(
            self, text=ico, text_color=col,
            font=ctk.CTkFont(size=13, weight="bold"), width=22
        )
        self._lbl_ico.pack(side="left", padx=(10, 4), pady=8)

        # Nombre + descripción de estado
        centro = ctk.CTkFrame(self, fg_color="transparent")
        centro.pack(side="left", fill="x", expand=True, pady=6)

        nombre = self._item["nombre"]
        nombre_corto = (nombre[:34] + "…") if len(nombre) > 35 else nombre
        self._lbl_nombre = ctk.CTkLabel(
            centro, text=nombre_corto,
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color=C["txt_main"], anchor="w"
        )
        self._lbl_nombre.pack(anchor="w")

        _, _, desc = _INFO_ESTADO.get(self._item["estado"], ("?", C["txt_dim"], ""))
        self._lbl_estado = ctk.CTkLabel(
            centro, text=desc,
            font=ctk.CTkFont(size=11),
            text_color=C["txt_dim"], anchor="w"
        )
        self._lbl_estado.pack(anchor="w")

        # Badge de campos con datos
        self._lbl_badge = ctk.CTkLabel(
            self, text="", font=ctk.CTkFont(size=11),
            text_color=C["txt_accent"], width=54, anchor="e"
        )
        self._lbl_badge.pack(side="right", padx=4)

        # Botón eliminar
        ctk.CTkButton(
            self, text="×", width=26, height=26, corner_radius=13,
            fg_color="transparent", hover_color=C["border"],
            text_color=C["txt_dim"], font=ctk.CTkFont(size=15),
            command=lambda: self._cb_del(self._item["ruta"])
        ).pack(side="right", padx=(0, 6))

        # Toda la fila es cliqueable para seleccionar
        self.bind("<Button-1>", lambda _e: self._cb_sel(self._item["ruta"]))
        for w in self.winfo_children():
            w.bind("<Button-1>", lambda _e: self._cb_sel(self._item["ruta"]))

        self.bind("<Enter>", lambda _: self._hover(True))
        self.bind("<Leave>", lambda _: self._hover(False))

    def _hover(self, dentro: bool):
        # Solo aplicar hover si no está seleccionado
        if self.cget("fg_color") != C["bg_sel"]:
            self.configure(fg_color=C["bg_row_alt"] if dentro else C["bg_card"])

    def actualizar(self):
        """Refresca icono, texto de estado y badge de campos."""
        ico, col, desc = _INFO_ESTADO.get(self._item["estado"], ("?", C["txt_dim"], ""))
        self._lbl_ico.configure(text=ico, text_color=col)
        self._lbl_estado.configure(text=desc)
        if self._item["metadata"]:
            n = sum(1 for v in self._item["metadata"].values() if v)
            self._lbl_badge.configure(text=f"{n} campos")

    def set_seleccionado(self, sel: bool):
        self.configure(fg_color=C["bg_sel"] if sel else C["bg_card"])


# ══════════════════════════════════════════════════════════════════════════
#  UI — Lista de archivos
# ══════════════════════════════════════════════════════════════════════════

class ListaArchivos(ctk.CTkScrollableFrame):
    def __init__(self, padre, cb_seleccionar, cb_eliminar, **kwargs):
        super().__init__(padre, **kwargs)
        self._filas: dict = {}       # ruta → FilaArchivo
        self._sel: str | None = None
        self._cb_sel = cb_seleccionar
        self._cb_del = cb_eliminar

    def agregar(self, item: dict):
        fila = FilaArchivo(self, item, self._seleccionar, self._cb_del)
        fila.pack(fill="x", pady=2)
        self._filas[item["ruta"]] = fila

    def actualizar_fila(self, ruta: str):
        if ruta in self._filas:
            self._filas[ruta].actualizar()

    def eliminar_fila(self, ruta: str):
        if ruta in self._filas:
            self._filas[ruta].destroy()
            del self._filas[ruta]
            if self._sel == ruta:
                self._sel = None

    def _seleccionar(self, ruta: str):
        if self._sel and self._sel in self._filas:
            self._filas[self._sel].set_seleccionado(False)
        self._sel = ruta
        if ruta in self._filas:
            self._filas[ruta].set_seleccionado(True)
        self._cb_sel(ruta)

    def seleccionar_primero(self):
        if self._filas:
            primera = next(iter(self._filas))
            self._seleccionar(primera)

    def vaciar(self):
        for f in self._filas.values():
            f.destroy()
        self._filas.clear()
        self._sel = None


# ══════════════════════════════════════════════════════════════════════════
#  UI — Zona de carga (drag & drop + clic)
# ══════════════════════════════════════════════════════════════════════════

class ZonaCarga(ctk.CTkFrame):
    """Acepta uno o varios archivos por arrastre o diálogo."""

    def __init__(self, padre, cb_archivos, **kwargs):
        super().__init__(padre, **kwargs)
        self._cb = cb_archivos   # cb(rutas: list[str])
        self._construir()
        self._activar_dnd()

    def _construir(self):
        self.configure(fg_color=C["bg_card"], border_color=C["border"],
                       border_width=2, corner_radius=14)

        ctk.CTkLabel(self, text="⬆", font=ctk.CTkFont(size=36),
                     text_color=C["accent"]).pack(pady=(18, 4))
        ctk.CTkLabel(self, text="Arrastra uno o varios archivos aquí",
                     font=ctk.CTkFont(size=14, weight="bold"),
                     text_color=C["txt_main"]).pack()
        ctk.CTkLabel(self, text="o haz clic para seleccionar",
                     font=ctk.CTkFont(size=12),
                     text_color=C["txt_dim"]).pack(pady=(2, 4))
        ctk.CTkLabel(self, text="JSON  ·  DOCX  ·  XLSX  ·  PDF  ·  PNG  ·  JPG",
                     font=ctk.CTkFont(size=11),
                     text_color=C["txt_accent"]).pack(pady=(0, 18))

        self.bind("<Button-1>", self._al_clic)
        for h in self.winfo_children():
            h.bind("<Button-1>", self._al_clic)

        self.bind("<Enter>", lambda _: self.configure(border_color=C["accent"]))
        self.bind("<Leave>", lambda _: self.configure(border_color=C["border"]))

    def _activar_dnd(self):
        if _DND_AVAILABLE:
            try:
                self.drop_target_register(DND_FILES)
                self.dnd_bind("<<Drop>>", self._al_soltar)
            except Exception:
                pass

    def _al_clic(self, _e=None):
        rutas = filedialog.askopenfilenames(
            title="Seleccionar archivos",
            filetypes=[
                ("Archivos soportados",
                 "*.json *.docx *.xlsx *.pdf *.png *.jpg *.jpeg"),
                ("Todos los archivos", "*.*"),
            ],
        )
        if rutas:
            self._cb(list(rutas))

    def _al_soltar(self, evento):
        rutas = _parsear_rutas_dnd(evento.data)
        if rutas:
            self._cb(rutas)


# ══════════════════════════════════════════════════════════════════════════
#  UI — Ventana de ayuda
# ══════════════════════════════════════════════════════════════════════════

def mostrar_ayuda(padre):
    win = ctk.CTkToplevel(padre)
    win.title("Como usar MetaCleaner")
    win.geometry("420x380")
    win.configure(fg_color=C["bg_panel"])
    win.grab_set()
    ctk.CTkLabel(win, text="Como usar MetaCleaner",
                 font=ctk.CTkFont(size=16, weight="bold"),
                 text_color=C["txt_main"]).pack(pady=(20, 12))
    pasos = [
        ("1", "Arrastra o selecciona uno o varios archivos."),
        ("2", "Haz clic en un archivo de la lista para ver sus metadatos."),
        ("3", "Pulsa \"Limpiar\" para eliminar los metadatos."),
        ("4", "Elige carpeta de destino (varios archivos) o nombre de salida (uno solo)."),
    ]
    for num, txt in pasos:
        f = ctk.CTkFrame(win, fg_color=C["bg_card"], corner_radius=8)
        f.pack(fill="x", padx=24, pady=3)
        ctk.CTkLabel(f, text=num, font=ctk.CTkFont(weight="bold"),
                     text_color=C["accent"], width=28).pack(side="left", padx=8, pady=9)
        ctk.CTkLabel(f, text=txt, text_color=C["txt_main"],
                     anchor="w").pack(side="left", padx=4)
    ctk.CTkLabel(win,
                 text="Formatos: JSON · DOCX · XLSX · PDF · PNG · JPG",
                 font=ctk.CTkFont(size=11), text_color=C["txt_accent"]
                 ).pack(pady=(14, 4))
    ctk.CTkButton(win, text="Cerrar", command=win.destroy,
                  fg_color=C["accent"], hover_color=C["accent_dim"],
                  corner_radius=10).pack(pady=14)


# ══════════════════════════════════════════════════════════════════════════
#  APLICACIÓN PRINCIPAL
# ══════════════════════════════════════════════════════════════════════════

if _DND_AVAILABLE:
    class _Base(ctk.CTk, _TkDnD.DnDWrapper):
        def __init__(self):
            super().__init__()
            self.TkdndVersion = _TkDnD._require(self)
else:
    class _Base(ctk.CTk):
        def __init__(self):
            super().__init__()


class MetaCleanerApp(_Base):

    def __init__(self):
        super().__init__()
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        # Cada item: {ruta, nombre, tamano, estado, metadata, msg}
        self._archivos: list = []

        self._configurar_ventana()
        self._construir_ui()

    # ── Ventana ──────────────────────────────────────────────────────────

    def _configurar_ventana(self):
        self.title("MetaCleaner — Limpiador de Metadatos")
        self.geometry("980x680")
        self.minsize(840, 580)
        self.configure(fg_color=C["bg_main"])
        self.update_idletasks()
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        self.geometry(f"980x680+{(sw-980)//2}+{(sh-680)//2}")

    # ── UI ───────────────────────────────────────────────────────────────

    def _construir_ui(self):
        raiz = ctk.CTkFrame(self, fg_color=C["bg_main"])
        raiz.pack(fill="both", expand=True, padx=20, pady=20)
        self._cabecera(raiz)
        self._cuerpo(raiz)
        self._pie(raiz)

    def _cabecera(self, p):
        bar = ctk.CTkFrame(p, fg_color=C["bg_panel"], corner_radius=12)
        bar.pack(fill="x", pady=(0, 12))
        ctk.CTkLabel(bar, text="MetaCleaner",
                     font=ctk.CTkFont(size=20, weight="bold"),
                     text_color=C["txt_main"]).pack(side="left", padx=20, pady=13)
        ctk.CTkLabel(bar, text="Privacidad  ·  Seguridad  ·  Control",
                     font=ctk.CTkFont(size=12),
                     text_color=C["txt_accent"]).pack(side="left", padx=4)
        ctk.CTkButton(bar, text="?", width=34, height=34, corner_radius=17,
                      fg_color=C["bg_card"], hover_color=C["border"],
                      text_color=C["txt_dim"], font=ctk.CTkFont(weight="bold"),
                      command=lambda: mostrar_ayuda(self)
                      ).pack(side="right", padx=16, pady=13)

    def _cuerpo(self, p):
        cont = ctk.CTkFrame(p, fg_color="transparent")
        cont.pack(fill="both", expand=True)
        cont.grid_columnconfigure(0, weight=2, minsize=290)
        cont.grid_columnconfigure(1, weight=3)
        cont.grid_rowconfigure(0, weight=1)

        izq = ctk.CTkFrame(cont, fg_color="transparent")
        izq.grid(row=0, column=0, sticky="nsew", padx=(0, 8))
        self._panel_izq(izq)

        der = ctk.CTkFrame(cont, fg_color="transparent")
        der.grid(row=0, column=1, sticky="nsew", padx=(8, 0))
        self._panel_der(der)

    def _panel_izq(self, p):
        # Encabezado con contador
        enc = ctk.CTkFrame(p, fg_color="transparent")
        enc.pack(fill="x", pady=(0, 6))
        ctk.CTkLabel(enc, text="Archivos",
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color=C["txt_dim"]).pack(side="left")
        self._lbl_contador = ctk.CTkLabel(enc, text="",
                                           font=ctk.CTkFont(size=11),
                                           text_color=C["txt_accent"])
        self._lbl_contador.pack(side="right")

        # Zona de carga (compacta)
        self._zona = ZonaCarga(p, self._al_agregar_archivos, height=140)
        self._zona.pack(fill="x")

        # Lista de archivos (scrollable)
        self._lista = ListaArchivos(
            p,
            cb_seleccionar=self._al_seleccionar_fila,
            cb_eliminar=self._al_eliminar_archivo,
            fg_color=C["bg_panel"],
            corner_radius=10,
        )
        self._lista.pack(fill="both", expand=True, pady=(8, 0))

        # Botón principal — Limpiar
        self._btn_limpiar = ctk.CTkButton(
            p, text="Limpiar archivos",
            font=ctk.CTkFont(size=14, weight="bold"),
            height=50, corner_radius=12,
            fg_color=C["accent"], hover_color=C["accent_dim"],
            state="disabled",
            command=self._al_limpiar,
        )
        self._btn_limpiar.pack(fill="x", pady=(8, 0))

        # Botón secundario — Vaciar lista
        self._btn_vaciar = ctk.CTkButton(
            p, text="Vaciar lista",
            font=ctk.CTkFont(size=12), height=34, corner_radius=10,
            fg_color="transparent", hover_color=C["bg_card"],
            border_color=C["border"], border_width=1,
            text_color=C["txt_dim"], state="disabled",
            command=self._vaciar_lista,
        )
        self._btn_vaciar.pack(fill="x", pady=(6, 0))

    def _panel_der(self, p):
        enc = ctk.CTkFrame(p, fg_color="transparent")
        enc.pack(fill="x", pady=(0, 8))
        ctk.CTkLabel(enc, text="Metadatos detectados",
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color=C["txt_dim"]).pack(side="left")
        self._lbl_archivo_activo = ctk.CTkLabel(enc, text="",
                                                 font=ctk.CTkFont(size=11),
                                                 text_color=C["txt_accent"])
        self._lbl_archivo_activo.pack(side="left", padx=8)
        self._lbl_conteo = ctk.CTkLabel(enc, text="",
                                         font=ctk.CTkFont(size=11),
                                         text_color=C["txt_dim"])
        self._lbl_conteo.pack(side="right")
        self._tabla = TablaMetadatos(p, fg_color=C["bg_panel"], corner_radius=12)
        self._tabla.pack(fill="both", expand=True)

    def _pie(self, p):
        pie = ctk.CTkFrame(p, fg_color=C["bg_panel"], corner_radius=10, height=50)
        pie.pack(fill="x", pady=(12, 0))
        pie.pack_propagate(False)
        self._progreso = ctk.CTkProgressBar(
            pie, mode="indeterminate",
            fg_color=C["bg_card"], progress_color=C["accent"]
        )
        self._lbl_estado = ctk.CTkLabel(
            pie, text="Listo  ·  Arrastra archivos para comenzar",
            font=ctk.CTkFont(size=12), text_color=C["txt_dim"]
        )
        self._lbl_estado.pack(side="left", padx=16, pady=13)

    # ── Callbacks ────────────────────────────────────────────────────────

    def _al_agregar_archivos(self, rutas: list):
        """Añade archivos a la lista, ignorando duplicados y no soportados."""
        nuevos = []
        ignorados = []
        rutas_existentes = {a["ruta"] for a in self._archivos}

        for ruta in rutas:
            ruta = ruta.strip()
            ext  = Path(ruta).suffix.lower()
            if ext not in MetadataExtractor.FORMATOS:
                ignorados.append(Path(ruta).name)
                continue
            if not os.path.isfile(ruta):
                continue
            if ruta in rutas_existentes:
                continue
            item = {
                "ruta":     ruta,
                "nombre":   os.path.basename(ruta),
                "tamano":   os.path.getsize(ruta),
                "estado":   "extrayendo",
                "metadata": {},
                "msg":      "",
            }
            self._archivos.append(item)
            self._lista.agregar(item)
            nuevos.append(item)

        if nuevos:
            self._actualizar_botones()
            # Si es el único archivo, seleccionarlo automáticamente
            if len(self._archivos) == len(nuevos):
                self._lista.seleccionar_primero()
            # Extraer metadatos de cada archivo nuevo en paralelo
            for item in nuevos:
                threading.Thread(target=self._hilo_extraer, args=(item,),
                                 daemon=True).start()

        if ignorados:
            self._estado(
                f"Ignorados (formato no soportado): {', '.join(ignorados[:3])}",
                "warning"
            )
        elif nuevos:
            n = len(nuevos)
            self._estado(f"{n} archivo{'s' if n > 1 else ''} cargado{'s' if n > 1 else ''}",
                         "info")

    def _al_seleccionar_fila(self, ruta: str):
        """Muestra los metadatos del archivo seleccionado."""
        item = next((a for a in self._archivos if a["ruta"] == ruta), None)
        if not item:
            return
        self._tabla.actualizar(item["metadata"])
        nombre_corto = (item["nombre"][:40] + "…") if len(item["nombre"]) > 41 else item["nombre"]
        self._lbl_archivo_activo.configure(text=nombre_corto)
        if item["metadata"]:
            n = sum(1 for v in item["metadata"].values() if v)
            sufijo = "  (limpio)" if item["estado"] == "limpio" else ""
            self._lbl_conteo.configure(text=f"{n} campos con datos{sufijo}")
        else:
            self._lbl_conteo.configure(text="")

    def _al_eliminar_archivo(self, ruta: str):
        self._archivos = [a for a in self._archivos if a["ruta"] != ruta]
        self._lista.eliminar_fila(ruta)
        self._actualizar_botones()
        if not self._archivos:
            self._tabla.limpiar()
            self._lbl_archivo_activo.configure(text="")
            self._lbl_conteo.configure(text="")
            self._estado("Listo  ·  Arrastra archivos para comenzar", "dim")

    def _al_limpiar(self):
        listos = [a for a in self._archivos if a["estado"] == "listo"]
        if not listos:
            return

        if len(listos) == 1:
            # Un solo archivo: diálogo "guardar como"
            item   = listos[0]
            origen = Path(item["ruta"])
            salida = filedialog.asksaveasfilename(
                title="Guardar archivo limpio",
                initialdir=str(origen.parent),
                initialfile=f"{origen.stem}_cleaned{origen.suffix}",
                defaultextension=origen.suffix,
                filetypes=[(f"Archivo {origen.suffix.upper()}", f"*{origen.suffix}"),
                           ("Todos los archivos", "*.*")],
            )
            if not salida:
                return
            destinos = {item["ruta"]: salida}
        else:
            # Varios archivos: elegir carpeta destino
            carpeta = filedialog.askdirectory(
                title=f"Carpeta de destino para {len(listos)} archivos limpios"
            )
            if not carpeta:
                return
            destinos = {
                a["ruta"]: str(
                    Path(carpeta) / f"{Path(a['ruta']).stem}_cleaned{Path(a['ruta']).suffix}"
                )
                for a in listos
            }

        self._btn_limpiar.configure(state="disabled")
        self._mostrar_progreso(True)
        threading.Thread(
            target=self._hilo_limpiar, args=(listos, destinos), daemon=True
        ).start()

    # ── Hilos de fondo ───────────────────────────────────────────────────

    def _hilo_extraer(self, item: dict):
        try:
            meta = MetadataExtractor.extraer(item["ruta"])
            self.after(0, self._extraccion_ok, item, meta)
        except Exception as e:
            self.after(0, self._extraccion_error, item, str(e))

    def _extraccion_ok(self, item: dict, meta: dict):
        item["estado"]   = "listo"
        item["metadata"] = meta
        self._lista.actualizar_fila(item["ruta"])
        self._actualizar_botones()
        # Si este archivo está seleccionado, actualizar la tabla
        if self._lista._sel == item["ruta"]:
            self._al_seleccionar_fila(item["ruta"])
        n_listos = sum(1 for a in self._archivos if a["estado"] == "listo")
        self._estado(f"{n_listos} archivo{'s' if n_listos > 1 else ''} listo{'s' if n_listos > 1 else ''} para limpiar",
                     "success")

    def _extraccion_error(self, item: dict, msg: str):
        item["estado"] = "error"
        item["msg"]    = msg
        self._lista.actualizar_fila(item["ruta"])
        self._actualizar_botones()
        self._estado(f"Error leyendo '{item['nombre']}': {msg}", "error")

    def _hilo_limpiar(self, archivos: list, destinos: dict):
        total = len(archivos)
        for i, item in enumerate(archivos, 1):
            self.after(0, self._limpieza_inicio, item, i, total)
            try:
                MetadataCleaner.limpiar(item["ruta"], destinos[item["ruta"]])
                self.after(0, self._limpieza_ok, item, destinos[item["ruta"]])
            except Exception as e:
                self.after(0, self._limpieza_error_item, item, str(e))
        self.after(0, self._limpieza_batch_fin, total)

    def _limpieza_inicio(self, item: dict, i: int, total: int):
        item["estado"] = "limpiando"
        self._lista.actualizar_fila(item["ruta"])
        self._estado(f"Limpiando {i} de {total}: '{item['nombre']}'…", "info")

    def _limpieza_ok(self, item: dict, ruta_salida: str):
        item["estado"] = "limpio"
        # Cargar metadatos del archivo ya limpio para verificación
        try:
            item["metadata"] = MetadataExtractor.extraer(ruta_salida)
        except Exception:
            pass
        self._lista.actualizar_fila(item["ruta"])
        if self._lista._sel == item["ruta"]:
            self._al_seleccionar_fila(item["ruta"])

    def _limpieza_error_item(self, item: dict, msg: str):
        item["estado"] = "error"
        item["msg"]    = msg
        self._lista.actualizar_fila(item["ruta"])

    def _limpieza_batch_fin(self, total: int):
        self._mostrar_progreso(False)
        limpios = sum(1 for a in self._archivos if a["estado"] == "limpio")
        errores = sum(1 for a in self._archivos if a["estado"] == "error")
        if errores == 0:
            self._estado(f"Completado  ·  {limpios} archivo{'s' if limpios > 1 else ''} limpio{'s' if limpios > 1 else ''}",
                         "success")
        else:
            self._estado(f"{limpios} limpios  ·  {errores} con error", "warning")
        self._actualizar_botones()

    # ── Utilidades ───────────────────────────────────────────────────────

    def _vaciar_lista(self):
        self._archivos.clear()
        self._lista.vaciar()
        self._tabla.limpiar()
        self._lbl_archivo_activo.configure(text="")
        self._lbl_conteo.configure(text="")
        self._actualizar_botones()
        self._mostrar_progreso(False)
        self._estado("Listo  ·  Arrastra archivos para comenzar", "dim")

    def _actualizar_botones(self):
        total   = len(self._archivos)
        listos  = sum(1 for a in self._archivos if a["estado"] == "listo")

        # Contador en el encabezado
        self._lbl_contador.configure(
            text=f"{total} archivo{'s' if total != 1 else ''}" if total else ""
        )

        # Botón limpiar
        if listos == 0:
            self._btn_limpiar.configure(state="disabled", text="Limpiar archivos")
        elif listos == 1:
            self._btn_limpiar.configure(state="normal", text="Limpiar 1 archivo",
                                        fg_color=C["accent"], hover_color=C["accent_dim"])
        else:
            self._btn_limpiar.configure(state="normal",
                                        text=f"Limpiar {listos} archivos",
                                        fg_color=C["accent"], hover_color=C["accent_dim"])

        # Botón vaciar
        self._btn_vaciar.configure(state="normal" if total > 0 else "disabled")

    def _estado(self, txt: str, nivel: str = "dim"):
        cols = {"dim": C["txt_dim"], "info": C["txt_main"],
                "success": C["success"], "error": C["error"], "warning": C["warning"]}
        self._lbl_estado.configure(text=txt, text_color=cols.get(nivel, C["txt_dim"]))

    def _mostrar_progreso(self, show: bool):
        if show:
            self._progreso.pack(side="right", padx=16, fill="x", expand=True)
            self._progreso.start()
        else:
            self._progreso.stop()
            self._progreso.pack_forget()


# ══════════════════════════════════════════════════════════════════════════
#  PUNTO DE ENTRADA
# ══════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    app = MetaCleanerApp()
    app.mainloop()
